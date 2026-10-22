"""Pliego de condiciones: estructura del anexo I (apartado 3) de la Orden ITC/1644/2011 y normas comprobadas."""

import json
from pathlib import Path

import docx

from ict_ajp.dimensionado import dimensionar
from ict_ajp.documento import Plantilla
from ict_ajp.modelo import Proyecto
from ict_ajp.pliego import NORMAS_GENERALES, generar_pliego

EJEMPLO = Path(__file__).parent.parent / "ejemplos" / "edificio_prueba.ict.json"


def test_estructura_y_contenido(tmp_path):
    p = Proyecto.cargar(EJEMPLO)
    ruta, titulos = generar_pliego(p, dimensionar(p), tmp_path / "p.docx", Plantilla.cargar(tmp_path))
    encabezados = [t for _, t in titulos]
    esperados = ["3.1.A", "3.1.B", "3.1.C", "3.1.D", "3.1.E", "3.1.F", "3.1.G",
                 "3.2.A", "3.2.B", "3.2.C", "3.2.D", "3.2.E", "3.2.F"]
    pos = [next(i for i, t in enumerate(encabezados) if t.startswith(e)) for e in esperados]
    assert pos == sorted(pos)
    d = docx.Document(str(ruta))
    texto = "\n".join([x.text for x in d.paragraphs] + [c.text for t in d.tables for f in t.rows for c in f.cells])
    assert "149922" in texto and "509812" in texto                      # productos del catálogo
    assert "≤ 20 dB entre el registro principal coaxial" in texto        # Anexo II 6.4.1
    assert "máxima 2 dB" in texto                                        # Anexo II 6.6.2
    assert "Ley 31/1995, de 8 de noviembre, de Prevención de Riesgos Laborales" in texto
    assert "Real Decreto 186/2016" in texto                              # sustituye al RD 1580/2006 derogado
    assert "Ley 11/2022, de 28 de junio, General de Telecomunicaciones" in texto


def test_reaccion_al_fuego_de_los_cables(tmp_path):
    from ict_ajp import normativa as N
    assert N.CLASE_FUEGO_CABLES_ICT == "Dca-s2,d2,a2"                    # Orden ECE/983/2019, anexo, tabla 1
    assert N.cumple_clase_fuego("Dca-s2,d2,a2") is True
    assert N.cumple_clase_fuego("Eca") is None                           # otra clase: la verifica el proyectista
    p = Proyecto.cargar(EJEMPLO)
    ruta, _ = generar_pliego(p, dimensionar(p), tmp_path / "p.docx", Plantilla.cargar(tmp_path))
    texto = "\n".join(x.text for x in docx.Document(str(ruta)).paragraphs)
    assert "clase de reacción al fuego Dca-s2,d2,a2" in texto and "414002) es de euroclase Dca-s2,d2,a2" in texto


def test_normas_generales_comprobadas_en_el_boe():
    normas = json.loads(NORMAS_GENERALES.read_text(encoding="utf-8"))["normas"]
    assert normas["rd_1580_2006"]["vigente"] is False                    # derogado: no se cita como vigente
    citadas = {"ley_31_1995", "rd_1627_1997", "rd_105_2008", "ley_11_2022", "rd_186_2016", "rd_314_2006"}
    assert all(normas[k]["vigente"] and normas[k]["identificador"].startswith("BOE-A-") for k in citadas)
