"""Memoria en Word: estructura del anexo I de la Orden ITC/1644/2011 y estilo AJP con plantilla de prueba."""

import datetime as dt
import json
from pathlib import Path

import docx
import pytest

from ict_ajp.dimensionado import dimensionar
from ict_ajp.documento import Plantilla
from ict_ajp.memoria import _dependencias, generar_memoria
from ict_ajp.modelo import Proyecto, Usuario

EJEMPLO = Path(__file__).parent.parent / "ejemplos" / "edificio_prueba.ict.json"


@pytest.fixture
def plantilla(tmp_path):
    datos = {"empresa": "Empresa de prueba", "direccion": "Calle Inventada 1", "cp_localidad": "00000 Pueblo",
             "localidad_firma": "Pueblo", "proyectista": {"nombre": "Proyectista de prueba", "titulacion": "Ingeniero",
                                                          "colegiado": "0", "colegio": "COLEGIO"}}
    (tmp_path / "empresa.json").write_text(json.dumps(datos), encoding="utf-8")
    return Plantilla.cargar(tmp_path)


def _texto(ruta):
    d = docx.Document(str(ruta))
    partes = [p.text for p in d.paragraphs]
    for t in d.tables:
        for f in t.rows:
            partes += [c.text for c in f.cells]
    return d, "\n".join(partes)


def test_estructura_del_anexo_i(tmp_path, plantilla):
    p = Proyecto.cargar(EJEMPLO)
    ruta, titulos = generar_memoria(p, dimensionar(p), tmp_path / "m.docx", plantilla, fecha=dt.date(2026, 10, 21))
    d, texto = _texto(ruta)
    # Orden ITC/1644/2011, anexo I, apartado 1: todos los epígrafes obligatorios, en orden
    esperados = ["1.1.A", "1.1.B", "1.1.C", "1.1.D", "1.2.A", "1.2.B", "1.2.C", "1.2.D", "1.2.E", "1.2.F"]
    encabezados = [t for _, t in titulos]
    posiciones = [next(i for i, t in enumerate(encabezados) if t.startswith(e)) for e in esperados]
    assert posiciones == sorted(posiciones)
    assert "Fdo.: Proyectista de prueba" in texto and "Pueblo, 21 de octubre de 2026" in texto
    assert "[A COMPLETAR" in texto or "[PENDIENTE" in texto        # lo que falta se marca, no se inventa
    # cabecera con los datos de la plantilla y pie con numeración de páginas
    cab = d.sections[0].header.tables[0].rows[0].cells[1].text
    assert "Calle Inventada 1" in cab
    assert "NUMPAGES" in d.sections[0].footer._element.xml


def test_sin_plantilla_no_falla(tmp_path):
    p = Proyecto.cargar(EJEMPLO)
    ruta, _ = generar_memoria(p, dimensionar(p), tmp_path / "m.docx", Plantilla.cargar(tmp_path / "no_existe"))
    assert ruta.exists()


def test_resumen_de_dependencias():
    u = Usuario.model_validate({"id": "1A", "estancias": [
        {"nombre": "SALON-COMEDOR", "tipo": "salon"}, {"nombre": "DORM. 1", "tipo": "dormitorio"},
        {"nombre": "DORM. 2", "tipo": "dormitorio"}, {"nombre": "BAÑO", "tipo": "bano"},
        {"nombre": "BAÑO", "tipo": "bano"}, {"nombre": "VESTIBULO", "tipo": "otra", "computa": False}]})
    assert _dependencias(u) == "salón-comedor, 2 dormitorios, 2 baños, vestíbulo"


def test_datos_de_los_planos_en_la_memoria(tmp_path, plantilla):
    p = Proyecto.cargar(EJEMPLO)
    _, sin = _texto(generar_memoria(p, dimensionar(p), tmp_path / "a.docx", plantilla)[0])
    assert "[PENDIENTE: ubicación del mástil]" in sin and "[PENDIENTE: medición de tubos" in sin
    p.planos.base_dxf = "base.dxf"
    p.planos.mastil_xy = (1.0, 2.0)
    p.planos.ubicacion_mastil = "la esquina trasera de la cubierta"
    p.planos.acceso_cubierta = "trampilla con escalera escamoteable"
    p.planos.materiales = {"Tubo Ø20 mm, canalización interior de usuario [m]": 312.46,
                           "Registro de toma RJ45 (pares trenzados) [ud]": 36.0}
    _, con = _texto(generar_memoria(p, dimensionar(p), tmp_path / "b.docx", plantilla)[0])
    assert "El mástil se sitúa en la esquina trasera de la cubierta" in con
    assert "Acceso a la cubierta: trampilla con escalera escamoteable." in con
    assert "El espacio reservado, de 1,5 x 1,5 m, está junto al mástil" in con
    assert "Su ubicación figura en los planos de planta." in con
    assert "Tubo Ø20 mm, canalización interior de usuario\n312,5\nm" in con           # cuadro de materiales
    assert "Registro de toma RJ45 (pares trenzados)\n36\nud" in con
    assert "[PENDIENTE: ubicación" not in con and "[PENDIENTE: medición de tubos" not in con
