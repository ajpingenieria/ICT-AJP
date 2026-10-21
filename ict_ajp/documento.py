"""Documentos Word con el estilo de AJP Ingeniería (memoria, pliego, presupuesto...).

El estilo y los datos de la empresa (logos, dirección, proyectista) se leen de la plantilla de Drive
(«_Plantilla ICT-AJP/empresa.json»), que no está en el repositorio. Sin plantilla se genera igual, sin
logos ni datos de contacto.

Marcas en el texto: «[PENDIENTE: ...]» y «[A VERIFICAR: ...]» salen resaltadas en amarillo para que el
proyectista las vea al revisar.
"""

from __future__ import annotations

import json
import os
import re
from dataclasses import dataclass, field
from pathlib import Path
from typing import Optional, Sequence

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_COLOR_INDEX, WD_TAB_ALIGNMENT, WD_TAB_LEADER
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

PLANTILLA_POR_DEFECTO = Path(r"G:\Mi unidad\Proyectos AJP\_Plantilla ICT-AJP")
_MARCA = re.compile(r"(\[(?:PENDIENTE|A VERIFICAR|A COMPLETAR)[^\]]*\])")


@dataclass
class Plantilla:
    carpeta: Optional[Path] = None
    datos: dict = field(default_factory=dict)

    @classmethod
    def cargar(cls, carpeta: Optional[str | Path] = None) -> "Plantilla":
        c = Path(carpeta or os.environ.get("ICT_AJP_PLANTILLA") or PLANTILLA_POR_DEFECTO)
        f = c / "empresa.json"
        if f.exists():
            return cls(c, json.loads(f.read_text(encoding="utf-8")))
        return cls(None, {"empresa": "AJP Ingeniería", "proyectista": {}})

    def estilo(self, clave: str, defecto):
        return self.datos.get("estilo", {}).get(clave, defecto)

    def imagen(self, clave: str) -> Optional[Path]:
        if self.carpeta and self.datos.get(clave):
            p = self.carpeta / self.datos[clave]
            return p if p.exists() else None
        return None


def _rgb(hexa: str) -> RGBColor:
    return RGBColor.from_string(hexa.upper())


def _campo(parrafo, codigo: str, run_modelo=None) -> None:
    """Campo de Word (PAGE, NUMPAGES...) dentro de un párrafo."""
    r = parrafo.add_run()
    if run_modelo is not None:
        r.font.name, r.font.size, r.font.color.rgb = run_modelo.font.name, run_modelo.font.size, run_modelo.font.color.rgb
    for tipo, texto in (("begin", None), (None, codigo), ("separate", None), (None, "1"), ("end", None)):
        if tipo:
            e = OxmlElement("w:fldChar")
            e.set(qn("w:fldCharType"), tipo)
        else:
            e = OxmlElement("w:instrText" if texto == codigo else "w:t")
            e.set(qn("xml:space"), "preserve")
            e.text = texto
        r._r.append(e)


def _sombrear(celda, hexa: str) -> None:
    tc = celda._tc.get_or_add_tcPr()
    sh = OxmlElement("w:shd")
    sh.set(qn("w:val"), "clear")
    sh.set(qn("w:color"), "auto")
    sh.set(qn("w:fill"), hexa)
    tc.append(sh)


def _bordes(tabla, color: str = "BFBFBF") -> None:
    tbl_pr = tabla._tbl.tblPr
    b = OxmlElement("w:tblBorders")
    for lado in ("top", "left", "bottom", "right", "insideH", "insideV"):
        e = OxmlElement(f"w:{lado}")
        e.set(qn("w:val"), "single")
        e.set(qn("w:sz"), "4")
        e.set(qn("w:color"), color)
        b.append(e)
    tbl_pr.append(b)


class DocumentoAJP:
    def __init__(self, plantilla: Plantilla, texto_pie: str):
        self.p = plantilla
        self.doc = Document()
        self.fuente = plantilla.estilo("fuente", "Arial")
        self.color = plantilla.estilo("color_texto", "1F1F1F")
        self.gris = plantilla.estilo("color_secundario", "555555")
        self.marca = plantilla.estilo("color_marca", "FBC213")
        self.titulos: list[tuple[int, str]] = []          # para el índice
        self._estilos()
        self._pagina(texto_pie)

    # --- configuración ---------------------------------------------------------
    def _estilos(self) -> None:
        st = self.doc.styles
        normal = st["Normal"]
        normal.font.name, normal.font.size, normal.font.color.rgb = self.fuente, Pt(10), _rgb(self.color)
        normal.element.rPr.rFonts.set(qn("w:eastAsia"), self.fuente)
        normal.paragraph_format.space_after = Pt(6)
        normal.paragraph_format.line_spacing = 1.15
        for nivel, tam in ((1, 11.5), (2, 10.5), (3, 10)):
            h = st[f"Heading {nivel}"]
            h.font.name, h.font.size, h.font.bold = self.fuente, Pt(tam), True
            h.font.italic = nivel == 3
            h.font.color.rgb = _rgb(self.color)
            h.paragraph_format.space_before = Pt(12 if nivel == 1 else 8)
            h.paragraph_format.space_after = Pt(6)
            h.paragraph_format.keep_with_next = True

    def _pagina(self, texto_pie: str) -> None:
        s = self.doc.sections[0]
        s.page_width, s.page_height = Cm(21), Cm(29.7)
        m = self.p.estilo("margenes_cm", {"izquierda": 2.7, "derecha": 2.7, "superior": 4.15, "inferior": 2.65})
        s.left_margin, s.right_margin = Cm(m["izquierda"]), Cm(m["derecha"])
        s.top_margin, s.bottom_margin = Cm(m["superior"]), Cm(m["inferior"])
        s.header_distance = Cm(self.p.estilo("cabecera_cm", 1.15))
        s.footer_distance = Cm(self.p.estilo("pie_cm", 0.92))
        ancho = (s.page_width - s.left_margin - s.right_margin)
        # Cabecera: logo | dirección
        t = s.header.add_table(1, 2, ancho)
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        logo = self.p.imagen("logo_cabecera")
        if logo:
            t.cell(0, 0).paragraphs[0].add_run().add_picture(str(logo), height=Cm(1.6))
        else:
            self._run(t.cell(0, 0).paragraphs[0], self.p.datos.get("empresa", ""), 12, bold=True)
        d = self.p.datos
        derecha = t.cell(0, 1).paragraphs[0]
        derecha.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        lineas = [d.get("direccion", ""), d.get("cp_localidad", ""),
                  " · ".join(x for x in (d.get("telefono", ""), d.get("email", "")) if x)]
        for k, linea in enumerate(x for x in lineas if x):
            r = self._run(derecha, linea, 7.5, color=self.gris)
            if k < len([x for x in lineas if x]) - 1:
                r.add_break()
        s.header.paragraphs[0].text = ""
        # Pie: logo pequeño, texto y «Página x de y»
        pie = s.footer.paragraphs[0]
        logo_pie = self.p.imagen("logo_pie")
        if logo_pie:
            pie.add_run().add_picture(str(logo_pie), height=Cm(0.7))
            pie.add_run("  ")
        modelo = self._run(pie, texto_pie, 7.5, color=self.gris)
        self._run(pie, "      Página ", 7.5, color=self.gris)
        _campo(pie, "PAGE", modelo)
        self._run(pie, " de ", 7.5, color=self.gris)
        _campo(pie, "NUMPAGES", modelo)

    def _run(self, parrafo, texto: str, tam: float = 10, bold: bool = False, color: Optional[str] = None,
             italic: bool = False):
        r = parrafo.add_run(texto)
        r.font.name, r.font.size, r.bold, r.italic = self.fuente, Pt(tam), bold, italic
        r.font.color.rgb = _rgb(color or self.color)
        return r

    # --- contenido --------------------------------------------------------------
    def portada(self, titulo: str, subtitulo: str, datos: Sequence[tuple[str, str]], fecha: str) -> None:
        for _ in range(4):
            self.doc.add_paragraph()
        logo = self.p.imagen("logo_cabecera")
        if logo:
            pl = self.doc.add_paragraph()
            pl.alignment = WD_ALIGN_PARAGRAPH.CENTER
            pl.add_run().add_picture(str(logo), width=Cm(8))
            self.doc.add_paragraph()
        p = self.doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        self._run(p, titulo.upper(), 16, bold=True)
        p = self.doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        self._texto_con_marcas(p, subtitulo, 11, color=self.gris)
        self.doc.add_paragraph()
        t = self.doc.add_table(rows=0, cols=2)
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        _bordes(t)
        for k, v in datos:
            fila = t.add_row().cells
            fila[0].width, fila[1].width = Cm(4.5), Cm(10)
            self._run(fila[0].paragraphs[0], k, 9.5, bold=True)
            _sombrear(fila[0], "F2F2F2")
            self._texto_con_marcas(fila[1].paragraphs[0], v, 9.5)
        self.doc.add_paragraph()
        p = self.doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        self._run(p, fecha, 10, color=self.gris)
        self.salto_pagina()

    def indice(self, entradas: Optional[list[tuple[int, str, Optional[int]]]] = None) -> None:
        """Índice estático: (nivel, título, página o None). Se rellena en una segunda pasada."""
        p = self.doc.add_paragraph()
        self._run(p, "ÍNDICE", 11.5, bold=True)
        for nivel, texto, pagina in entradas or []:
            q = self.doc.add_paragraph()
            q.paragraph_format.left_indent = Cm(0.6 * (nivel - 1))
            q.paragraph_format.space_after = Pt(1)
            ancho = self.doc.sections[0].page_width - self.doc.sections[0].left_margin - self.doc.sections[0].right_margin
            q.paragraph_format.tab_stops.add_tab_stop(ancho, WD_TAB_ALIGNMENT.RIGHT, WD_TAB_LEADER.DOTS)
            self._run(q, texto, 9.5 if nivel > 1 else 10, bold=nivel == 1)
            self._run(q, "\t" + (str(pagina) if pagina else ""), 9.5)
        self.salto_pagina()

    def titulo(self, texto: str, nivel: int = 1) -> None:
        self.titulos.append((nivel, texto))
        h = self.doc.add_heading(texto.upper() if nivel == 1 else texto, level=nivel)
        for r in h.runs:
            r.font.name, r.font.color.rgb = self.fuente, _rgb(self.color)

    def _texto_con_marcas(self, p, texto: str, tam: float = 10, bold: bool = False, color: Optional[str] = None):
        for trozo in _MARCA.split(texto):
            if not trozo:
                continue
            r = self._run(p, trozo, tam, bold=bold, color=color)
            if _MARCA.fullmatch(trozo):
                r.font.highlight_color = WD_COLOR_INDEX.YELLOW
                r.bold = True

    def parrafo(self, texto: str, justificado: bool = True, tam: float = 10, bold: bool = False,
                color: Optional[str] = None) -> None:
        p = self.doc.add_paragraph()
        if justificado:
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        if bold or texto.rstrip().endswith(":"):     # subtítulos y entradillas: con lo que sigue
            p.paragraph_format.keep_with_next = True
        self._texto_con_marcas(p, texto, tam, bold, color)

    def vinetas(self, items: Sequence[str]) -> None:
        for it in items:
            p = self.doc.add_paragraph(style="List Bullet")
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            p.paragraph_format.space_after = Pt(2)
            self._texto_con_marcas(p, it)

    def nota(self, texto: str) -> None:
        self.parrafo(texto, tam=8.5, color=self.gris)

    def tabla(self, cabeceras: Sequence[str], filas: Sequence[Sequence], anchos_cm: Optional[Sequence[float]] = None,
              col_cumple: Optional[int] = None, tam: float = 8.5) -> None:
        t = self.doc.add_table(rows=1, cols=len(cabeceras))
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        t.autofit = False
        _bordes(t)
        _repetir_cabecera(t.rows[0])
        for j, c in enumerate(cabeceras):
            celda = t.rows[0].cells[j]
            _sombrear(celda, "1F1F1F")
            self._run(celda.paragraphs[0], str(c), tam, bold=True, color="FFFFFF")
        for fila in filas:
            celdas = t.add_row().cells
            for j, v in enumerate(fila):
                txt = _formato(v)
                color = None
                if col_cumple is not None and j == col_cumple:
                    color = "1E7B34" if txt.lower().startswith("cumple") else ("C00000" if "no" in txt.lower() else None)
                self._texto_con_marcas(celdas[j].paragraphs[0], txt, tam, bold=color is not None, color=color)
        ancho_util = (self.doc.sections[0].page_width - self.doc.sections[0].left_margin
                      - self.doc.sections[0].right_margin) / 360000          # cm
        anchos = list(anchos_cm) if anchos_cm else _anchos_automaticos(cabeceras, filas, ancho_util)
        _fijar_anchos(t, [Cm(w) for w in anchos])
        for k, fila in enumerate(t.rows):
            _no_partir(fila)
            if len(t.rows) <= 15 and k < len(t.rows) - 1:     # tablas cortas: enteras en una página
                for c in fila.cells:
                    for par in c.paragraphs:
                        par.paragraph_format.keep_with_next = True
        self.doc.add_paragraph().paragraph_format.space_after = Pt(2)

    def firma(self, fecha_texto: str) -> None:
        d = self.p.datos
        pr = d.get("proyectista", {})
        self.doc.add_paragraph()
        self.parrafo(f"{d.get('localidad_firma', '')}, {fecha_texto}".strip(", "), justificado=False)
        for _ in range(3):
            self.doc.add_paragraph()
        p = self.doc.add_paragraph()
        self._run(p, f"Fdo.: {pr.get('nombre', '[A COMPLETAR: proyectista]')}", 10, bold=True)
        lin = " · ".join(x for x in (pr.get("titulacion"), f"Colegiado n.º {pr['colegiado']}" if pr.get("colegiado")
                                     else None, pr.get("colegio")) if x)
        if lin:
            self.parrafo(lin, justificado=False, tam=8.5, color=self.gris)

    def salto_pagina(self) -> None:
        self.doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

    def guardar(self, ruta: str | Path) -> Path:
        ruta = Path(ruta)
        self.doc.save(ruta)
        return ruta


def _repetir_cabecera(fila) -> None:
    tr_pr = fila._tr.get_or_add_trPr()
    e = OxmlElement("w:tblHeader")
    e.set(qn("w:val"), "true")
    tr_pr.append(e)


def _no_partir(fila) -> None:
    tr_pr = fila._tr.get_or_add_trPr()
    e = OxmlElement("w:cantSplit")
    e.set(qn("w:val"), "true")
    tr_pr.append(e)


def _fijar_anchos(tabla, anchos) -> None:
    """Anchos en la rejilla de la tabla y en cada celda (Word y LibreOffice los respetan así)."""
    grid = tabla._tbl.tblGrid
    for j, col in enumerate(grid.findall(qn("w:gridCol"))):
        col.set(qn("w:w"), str(int(anchos[j].twips if hasattr(anchos[j], "twips") else anchos[j] / 635)))
    for fila in tabla.rows:
        for j, c in enumerate(fila.cells):
            c.width = anchos[j]


def _anchos_automaticos(cabeceras, filas, total_cm: float) -> list[float]:
    """Reparte el ancho según la longitud del texto de cada columna (mínimo 1,4 cm)."""
    pesos, minimos = [], []
    for j, c in enumerate(cabeceras):
        largos = [len(str(c))] + [len(_formato(f[j])) for f in filas if j < len(f)]
        pesos.append(max(4.0, min(40.0, sum(sorted(largos)[-3:]) / 3)))
        # que la palabra más larga de la cabecera quepa sin partirse (≈ 0,19 cm por letra a 8,5 pt)
        minimos.append(max(1.4, 0.19 * max(len(w) for w in str(c).split() or [""]) + 0.3))
    total = sum(pesos)
    anchos = [max(m, total_cm * p / total) for p, m in zip(pesos, minimos)]
    exceso = sum(anchos) - total_cm
    if exceso > 0:
        mayor = anchos.index(max(anchos))
        anchos[mayor] -= exceso
    return anchos


def _formato(v) -> str:
    """Números con coma decimal; el resto tal cual."""
    if isinstance(v, bool):
        return "Sí" if v else "No"
    if isinstance(v, float):
        s = f"{v:.2f}".rstrip("0").rstrip(".")
        return s.replace(".", ",")
    if isinstance(v, int) and abs(v) >= 10000:
        return f"{v:,}".replace(",", ".")
    return "" if v is None else str(v)
